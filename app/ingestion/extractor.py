from pathlib import Path
import json, csv, shutil, xml.etree.ElementTree as ET
SUPPORTED={".pdf",".docx",".txt",".rtf",".odt",".jpg",".jpeg",".png",".tif",".tiff",".bmp",".webp",
           ".xlsx",".xlsm",".ods",".csv",".pptx",".json",".xml",".yaml",".yml",".eml"}

def _get_tesseract():
    from app.core.config import settings
    cmd = settings.tesseract_cmd
    if cmd and Path(cmd).is_file():
        import pytesseract
        pytesseract.pytesseract.tesseract_cmd = cmd
        return pytesseract
    if shutil.which("tesseract"):
        import pytesseract
        return pytesseract
    return None

def extract(path):
    ext=Path(path).suffix.lower(); blocks=[]
    if ext not in SUPPORTED: raise ValueError(f"Unsupported extension: {ext}")
    if ext==".pdf":
        import pymupdf
        with pymupdf.open(path) as pdf:
            for i,p in enumerate(pdf,1):
                text=p.get_text().strip()
                if not text:
                    pytess = _get_tesseract()
                    if pytess:
                        try:
                            from PIL import Image
                            from io import BytesIO
                            text=pytess.image_to_string(Image.open(BytesIO(p.get_pixmap().tobytes("png")))).strip()
                        except Exception:
                            pass
                if text: blocks.append((f"page {i}",text))
    elif ext==".docx":
        from docx import Document
        d=Document(path); text="\n".join(p.text for p in d.paragraphs if p.text.strip())
        for ti,t in enumerate(d.tables,1):
            text += "\n" + "\n".join(f"[Table {ti}] "+" | ".join(c.text for c in row.cells) for row in t.rows)
        blocks=[("document",text)]
    elif ext in {".jpg",".jpeg",".png",".tif",".tiff",".bmp",".webp"}:
        pytess = _get_tesseract()
        if not pytess:
            raise RuntimeError(
                f"Tesseract OCR is not installed or configured. Cannot extract text from '{Path(path).name}'. "
                f"Install Tesseract OCR or set TESSERACT_CMD in .env to the tesseract.exe path."
            )
        from PIL import Image, ImageOps
        blocks=[("image",pytess.image_to_string(ImageOps.grayscale(Image.open(path))))]
    elif ext in {".txt",".yaml",".yml",".rtf"}:
        text=Path(path).read_text(encoding="utf-8",errors="replace")
        if ext==".rtf":
            from striprtf.striprtf import rtf_to_text
            text=rtf_to_text(text)
        blocks=[("document",text)]
    elif ext in {".xlsx",".xlsm",".ods"}:
        import pandas as pd
        engine="odf" if ext==".ods" else None
        x=pd.ExcelFile(path,engine=engine)
        for sheet in x.sheet_names:
            df=pd.read_excel(path,sheet_name=sheet,dtype=str,engine=engine).fillna("")
            text="Sheet: "+sheet+"\nColumns: "+" | ".join(map(str,df.columns))+"\n"
            for idx,row in df.iterrows(): text+=f"Row {idx+2}: "+" | ".join(f"{c}={row[c]}" for c in df.columns)+"\n"
            blocks.append((f"sheet {sheet}",text))
    elif ext==".csv":
        with open(path,encoding="utf-8-sig",newline="") as f: rows=list(csv.reader(f))
        if rows: blocks=[("CSV","\n".join(" | ".join(row) for row in rows))]
    elif ext==".pptx":
        from pptx import Presentation
        for i,s in enumerate(Presentation(path).slides,1):
            text="\n".join(sh.text for sh in s.shapes if hasattr(sh,"text") and sh.text.strip())
            if text: blocks.append((f"slide {i}",text))
    elif ext==".json":
        blocks=[("JSON",json.dumps(json.loads(Path(path).read_text(encoding="utf-8")),ensure_ascii=False,indent=2))]
    elif ext==".xml":
        blocks=[("XML",ET.tostring(ET.parse(path).getroot(),encoding="unicode"))]
    elif ext==".eml":
        from email import policy
        from email.parser import BytesParser
        m=BytesParser(policy=policy.default).parsebytes(Path(path).read_bytes()); body=m.get_body(preferencelist=("plain",))
        blocks=[("email",f"Subject: {m.get('subject','')}\nFrom: {m.get('from','')}\n"+(body.get_content() if body else ""))]
    elif ext==".odt":
        import zipfile
        with zipfile.ZipFile(path) as z: xml=z.read("content.xml")
        tree=ET.fromstring(xml); ns={"t":"urn:oasis:names:tc:opendocument:xmlns:text:1.0"}
        blocks=[("ODT","\n".join("".join(p.itertext()) for p in tree.findall(".//t:p",ns)))]
    blocks=[(loc,text) for loc,text in blocks if text.strip()]
    if not blocks: raise ValueError("No readable text extracted; document may be empty or OCR failed.")
    return blocks
